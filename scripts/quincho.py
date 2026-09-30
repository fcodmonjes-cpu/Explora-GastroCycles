#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
quincho.py — RESPALDO. Desde el 2026-09-30 el quincho se arma, publica e
imprime desde la página (PGO → Quincho, ARCHITECTURE §3.5), con el mismo
cálculo portado a JavaScript. Este script queda para trabajar sin la app o
desde la línea de comandos; si las dos versiones difieren, manda la página.

Reparte a los viajeros en las mesas del quincho y produce lo que el servicio
necesita: la planilla, los carteles y la publicación en el Handbook.

El criterio es el del quincho del 21-09-2026: se sientan juntos los que
compartieron exploraciones — misma exploración, mismo turno, mismo día —, y
ahora pesan más las recientes. Lo que ese día se decidió a mano pasa a ser
regla: los grupos que no caben en una mesa se llevan mesas unidas, las familias
pueden quedar con mesa propia, y las mesas sin vínculo se arman por idioma y
edad.

EL SCRIPT PROPONE; LA PLANILLA MANDA. Todo lo que una persona cambia en la
planilla se respeta al volver a calcular: si alguien movió a un viajero de mesa
(o lo marcó en la columna Fijo), ese viajero queda donde lo dejaron y el
cálculo reparte sólo al resto. Lo mismo con el Salón: se suman mesas, se cambia
cuántos caben o se reserva una para un grupo, y el cálculo trabaja sobre eso.

Pasos de un quincho:

  1. python scripts/quincho.py proponer --fecha 2026-10-05
       Lee los viajeros de /viajeros/current (el sync de PGO de ese día) y
       escribe quincho-2026-10-05.xlsx + quincho-2026-10-05-carteles.html.
  2. Se revisa la planilla en Excel y se ajusta lo que haga falta.
  3. python scripts/quincho.py rehacer quincho-2026-10-05.xlsx
       Vuelve a calcular respetando lo ajustado, y regenera los carteles.
       (carteles ARCHIVO regenera sólo los carteles, sin recalcular.)
  4. Carteles: abrir el HTML en Chrome, corregir con un clic si hace falta,
     elegir el color del logo e imprimir. Papel A4, se dobla al medio.
  5. python scripts/quincho.py publicar quincho-2026-10-05.xlsx
       La sube al Handbook (/quincho/current): Viajeros muestra la vista
       Quincho y cualquiera puede mover a un viajero de mesa desde el teléfono.
       --debug muestra lo que subiría sin escribir.
  6. python scripts/quincho.py bajar quincho-2026-10-05.xlsx
       Trae a la planilla los cambios hechos en el Handbook (quedan como Fijo).

Quien no usa el Handbook no pierde nada: la planilla y los carteles son
completos por sí solos.

Opciones de proponer:
  --fecha YYYY-MM-DD      noche del quincho (por defecto, hoy)
  --roster doc.json       lee el roster de un archivo en vez de Firebase (el que
                          deja `sync_viajeros.py --from-pgo --emit-json doc.json`)
  --mesas N --puestos P   el salón: N mesas de P puestos en fila (12 de 6)
  --puntas K              puestos extra en las cabeceras de una mesa unida (2)
  --sumar-mesas N         mesas extra al final del salón
  --reservar GRUPO:K      bloquea K mesas unidas para GRUPO (repetible)
  --mesa-propia GRUPO     el grupo va solo en su mesa (repetible)
  --excluir TEXTO         hab o parte del nombre de quien no va (repetible)
  --salida DIR            dónde escribir (por defecto, la carpeta actual)

Credenciales de Firebase: FIREBASE_KEY (service account, igual que el sync) o,
si no está, la clave del salón de 4 dígitos (ATA_PIN o se pregunta al correr).

PRIVACIDAD: la planilla y los carteles llevan nombres completos con su
habitación. El .gitignore los deja fuera del repo; no los subas a mano. El doc
del Handbook lo borra la propia app a los pocos días del evento.
"""

import argparse, datetime, getpass, hashlib, html, json, os, random, re, shutil, sys, base64
from itertools import combinations

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from sync_viajeros import DB_URL, norm_key, pid_de, parse_historia  # noqa: E402

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
QUINCHO_PATH = "quincho/current"

# ── El salón por defecto ──────────────────────────────────────────────────────
# Doce mesas de seis en una fila, como el 21-09. Una mesa unida (dos o más
# juntas) gana además las dos cabeceras: así 2×6 + 2 = 14, lo de Patil y
# Kangaroo.
MESAS_BASE   = 12
PUESTOS_BASE = 6
PUNTAS       = 2

# ── Pesos del cálculo ─────────────────────────────────────────────────────────
# Una exploración compartida vale 1 + 2·½^días: hoy 3, ayer 2, anteayer 1,5, y
# las viejas tienden a 1. "Cuentan todas, sobre todo las más recientes".
def peso_reciente(dias):
    return 1 + 2 * 0.5 ** dias

# Misma exploración y mismo turno pero en otro vehículo ("Astronomia" vs
# "Astronomia-2"): estuvieron en el mismo lugar a la misma hora, pero no
# necesariamente juntos. Cuenta, con menos peso.
PESO_PARCIAL  = 0.35
# Idioma y edad DESEMPATAN, no compiten: el 21-09 se usaron sólo para las mesas
# sin vínculo. Con pesos más altos (0,6 y 0,25 en la primera versión) el cálculo
# cambiaba exploraciones compartidas por afinidad de idioma, y su distribución
# salía peor que la hecha a mano. Por eso los dos juntos, en una mesa, no
# alcanzan a una exploración compartida.
BONO_IDIOMA   = 0.15
# Por cada 10 años de diferencia de edad promedio entre dos unidades.
CASTIGO_EDAD  = 0.05
# Escala de la raíz de afinidad de cada unidad con su mesa (ver score_mesa).
PESO_VINCULO  = 1.5
# Al RECALCULAR, quedarse en la mesa que ya tenía. Con los carteles impresos,
# un ajuste a mano no puede reordenar el salón entero: el cálculo mueve a los
# desplazados y a quien gane claramente con el cambio, nada más. `rehacer
# --desde-cero` lo ignora.
ESTABILIDAD   = 3.0
# Cada mesa en uso: empuja a llenar mesas antes que abrir una nueva.
CASTIGO_MESA  = 4.0
# Puestos vacíos al cuadrado en una mesa en uso: dos vacíos repartidos en dos
# mesas (1+1) pesan menos que dos en la misma (4).
CASTIGO_VACIO = 0.5

# nac de PGO → idioma de conversación. Lo que no está acá cuenta como inglés,
# que es la lengua común de la mesa (el 21-09 los suizos y el holandés se
# sentaron con los ingleses).
IDIOMA = {
    **{c: "es" for c in ("ARGE", "CHIL", "MEXI", "SPAI", "ESPA", "PARA", "ECUA", "PERU",
                         "SALV", "SV", "COLO", "URUG", "VENE", "BOLI", "COST", "GUAT",
                         "HOND", "NICA", "PANA", "DOMI", "CUBA", "PUER")},
    **{c: "pt" for c in ("BRAZ", "BRA", "PORT")},
    **{c: "fr" for c in ("FREN", "FRAN")},
    **{c: "it" for c in ("ITAL",)},
    **{c: "de" for c in ("GERM",)},
}
IDIOMA_NOMBRE = {"es": "español", "pt": "portugués", "en": "inglés",
                 "fr": "francés", "it": "italiano", "de": "alemán"}

SI = {"x", "si", "sí", "s", "1", "yes", "y", "✓", "fijo"}
NO_VA = {"-", "no", "no va"}


def aviso(msg):
    print(f"[quincho] {msg}")


def falla(msg):
    print(f"[quincho] ✗ {msg}")
    sys.exit(1)


def orden_hab(h):
    """'07' < '11-12' < '21' < 'SIN HAB': número primero, texto después."""
    m = re.match(r"\d+", str(h))
    return (0, int(m.group()), str(h)) if m else (1, 0, str(h))


# ── Firebase ──────────────────────────────────────────────────────────────────
# Dos llaves, la misma puerta. La service account es la del sync (Actions); la
# clave del salón es la que teclea el equipo, derivada igual que en la app
# (ataPinPassword en index.html): PBKDF2-SHA256, 200.000 vueltas, esa sal.
FB_API_KEY = "AIzaSyDVyaPMAPVsbQFUhsGTYnnOdQQH1zffjfc"   # CF_FIREBASE_CONFIG, pública por diseño
GATE_EMAIL = "salon@ata-handbook.local"
GATE_SALT  = "ata-handbook/salon/2026"
GATE_ITER  = 200000

_CRED = None


def fb_credencial():
    global _CRED
    if _CRED:
        return _CRED
    import requests
    if os.environ.get("FIREBASE_KEY"):
        from google.oauth2 import service_account
        from google.auth.transport.requests import Request as GoogleAuthRequest
        creds = service_account.Credentials.from_service_account_info(
            json.loads(os.environ["FIREBASE_KEY"]),
            scopes=["https://www.googleapis.com/auth/firebase.database",
                    "https://www.googleapis.com/auth/userinfo.email"])
        creds.refresh(GoogleAuthRequest())
        _CRED = ({"Authorization": f"Bearer {creds.token}"}, {})
        return _CRED
    pin = os.environ.get("ATA_PIN") or getpass.getpass("Clave del salón (4 dígitos): ")
    pwd = hashlib.pbkdf2_hmac("sha256", pin.strip().encode(), GATE_SALT.encode(), GATE_ITER, 32).hex()
    r = requests.post(
        f"https://identitytoolkit.googleapis.com/v1/accounts:signInWithPassword?key={FB_API_KEY}",
        json={"email": GATE_EMAIL, "password": pwd, "returnSecureToken": True}, timeout=20)
    if not r.ok:
        falla("Firebase no aceptó la clave del salón.")
    _CRED = ({}, {"auth": r.json()["idToken"]})
    return _CRED


def fb(metodo, path, data=None):
    import requests
    headers, params = fb_credencial()
    r = requests.request(metodo, f"{DB_URL}/{path}.json", headers=headers, params=params,
                         data=None if data is None else json.dumps(data, ensure_ascii=False).encode("utf-8"),
                         timeout=30)
    if not r.ok:
        falla(f"Firebase respondió {r.status_code} en {metodo} /{path}. "
              "¿Está la regla de /quincho publicada en la consola?" if path.startswith("quincho")
              else f"Firebase respondió {r.status_code} en {metodo} /{path}.")
    return r.json()


# ── Viajeros ──────────────────────────────────────────────────────────────────
def fecha_de_ddmm(ddmm, ref):
    """'20-09' → date. PGO no da el año: si cae después de ref, es del año pasado."""
    try:
        d, m = (int(x) for x in str(ddmm).split("-")[:2])
        f = datetime.date(ref.year, m, d)
    except (ValueError, TypeError):
        return None
    return f if f <= ref else f.replace(year=ref.year - 1)


def exps_de(tr, fecha, fecha_doc):
    """Exploraciones del viajero como [(nombre, turno, 'YYYY-MM-DD')], sin HOTEL."""
    items = []
    hist = tr.get("historia") or []
    if not hist and tr.get("historiaTxt"):
        hist, _ = parse_historia(tr["historiaTxt"])
    for h in hist:
        f = fecha_de_ddmm(h.get("d"), fecha)
        if f:
            items.append((h.get("n") or "", (h.get("t") or "").upper(), f.isoformat()))
    # La exploración de HOY (columna exp del Geos) es del día del reporte.
    exp = tr.get("exp") or []
    if isinstance(exp, dict):
        exp = [{"t": exp.get("turno"), "n": exp.get("txt")}]
    for e in exp:
        if e and e.get("n"):
            items.append((e["n"], (e.get("t") or "").upper(), fecha_doc.isoformat()))
    fuera = []
    for n, t, d in items:
        if norm_key(n).upper().startswith("HOTEL"):
            continue
        dd = datetime.date.fromisoformat(d)
        if dd > fecha:
            continue
        # Una NOC del mismo día del quincho es DURANTE el quincho: no la compartieron antes.
        if dd == fecha and t == "NOC":
            continue
        fuera.append((n, t, d))
    return sorted(set(fuera))


def personas_del_roster(doc, fecha, excluir):
    """Quienes duermen en el hotel la noche del quincho: in ≤ fecha < out."""
    fecha_doc = datetime.date.fromisoformat(doc.get("date") or fecha.isoformat())
    if fecha_doc != fecha:
        aviso(f"El roster es del {fecha_doc} y el quincho del {fecha}. Los que todavía no "
              "llegaron al hotel no están en el roster: se suman a mano en la planilla.")
    fi = fecha.isoformat()
    out, fuera = [], 0
    for hab, lista in (doc.get("habs") or {}).items():
        for tr in lista or []:
            if (tr.get("in") or "0000") > fi or (tr.get("out") or "9999") <= fi:
                fuera += 1
                continue
            p = {"pid": tr.get("pid") or pid_de(tr.get("nombre")), "nombre": tr.get("nombre") or "",
                 "hab": str(hab), "edad": tr.get("edad"), "nac": tr.get("nac") or "",
                 "grupo": (tr.get("grupo") or "").strip(), "exps": exps_de(tr, fecha, fecha_doc),
                 "mesa": None, "fijo": False}
            txt = norm_key(p["nombre"])
            if any(norm_key(x) == norm_key(p["hab"]) or norm_key(x) in txt for x in excluir):
                p["mesa"] = "-"
            out.append(p)
    if fuera:
        aviso(f"{fuera} viajero(s) del roster no duermen en el hotel esa noche: quedan fuera.")
    return out


def clave_unidad(p):
    return (p["grupo"] or f"HAB {p['hab']}").upper()


def nombre_unidad(clave, personas):
    """'LOMBARDI' → 'Lombardi'. Sin grupo, el apellido del primero."""
    if clave.startswith("HAB "):
        partes = (personas[0]["nombre"] or "").split()
        return partes[-1].title() if partes else clave
    return clave.title()


# ── Unidades y afinidad ───────────────────────────────────────────────────────
class Unidad:
    def __init__(self, clave, personas, fecha):
        self.clave = clave
        self.personas = personas
        self.n = len(personas)
        self.nombre = nombre_unidad(clave, personas)
        self.exps, self.bases = {}, {}
        for p in personas:
            for n, t, d in p["exps"]:
                dias = (fecha - datetime.date.fromisoformat(d)).days
                w = peso_reciente(dias)
                k = (norm_key(n), t, d)
                self.exps[k] = max(self.exps.get(k, 0), w)
                b = (re.sub(r"-\d+$", "", norm_key(n)).strip(), t, d)
                self.bases[b] = max(self.bases.get(b, 0), w)
        idiomas = [IDIOMA.get(str(p["nac"]).upper(), "en") for p in personas]
        self.idioma = max(set(idiomas), key=idiomas.count) if idiomas else "en"
        edades = [p["edad"] for p in personas if isinstance(p["edad"], (int, float))]
        self.edad = sum(edades) / len(edades) if edades else None
        self.edades = edades


def afinidad(a, b):
    """Sólo exploraciones: el puntaje que se reporta y que arma las notas."""
    s = sum(w for k, w in a.exps.items() if k in b.exps)
    exactas = {(re.sub(r"-\d+$", "", k[0]).strip(), k[1], k[2]) for k in a.exps if k in b.exps}
    s += PESO_PARCIAL * sum(w for k, w in a.bases.items() if k in b.bases and k not in exactas)
    return s


def secundario(a, b):
    """Idioma y edad: lo que desempata entre mesas sin exploración compartida."""
    s = 0.0
    if a.idioma == b.idioma:
        s += BONO_IDIOMA
    if a.edad is not None and b.edad is not None:
        s -= CASTIGO_EDAD * abs(a.edad - b.edad) / 10
    return s


# ── Salón ─────────────────────────────────────────────────────────────────────
def mesas_para(n, puestos, puntas):
    """Cuántas mesas unidas necesita un grupo de n."""
    if n <= puestos:
        return 1
    k = 2
    while n > k * puestos + puntas:
        k += 1
    return k


def armar_salon(n_fisicas, puestos, puntas, reservas):
    """Mesas físicas en fila → mesas lógicas, con los bloques unidos para los grupos.

    reservas: {clave: k mesas}. Los bloques se reparten a lo largo del salón
    (el 21-09 los dos grupos grandes quedaron 4º y 8º de 10): cada uno busca el
    hueco libre más cercano a su lugar ideal. Las mesas se numeran de izquierda
    a derecha, así el número dice dónde está.
    """
    libres = [True] * n_fisicas
    bloques = {}
    orden = sorted(reservas.items(), key=lambda kv: -kv[1])
    for i, (clave, k) in enumerate(orden):
        ideal = (i + 1) / (len(orden) + 1) * n_fisicas
        cands = [s for s in range(n_fisicas - k + 1) if all(libres[s:s + k])]
        if not cands:
            falla(f"No hay {k} mesas seguidas libres para {clave}. Suma mesas con --sumar-mesas.")
        s = min(cands, key=lambda s: abs(s + k / 2 - ideal))
        for j in range(s, s + k):
            libres[j] = False
        bloques[s] = (clave, k)
    mesas, j = [], 0
    while j < n_fisicas:
        if j in bloques:
            clave, k = bloques[j]
            mesas.append({"puestos": k * puestos + (puntas if k > 1 else 0), "fisicas": k,
                          "reservada": clave, "nota": ""})
            j += k
        else:
            mesas.append({"puestos": puestos, "fisicas": 1, "reservada": "", "nota": ""})
            j += 1
    for i, m in enumerate(mesas, 1):
        m["n"] = i
        m["pos"] = i
    return mesas


# ── Optimizador ───────────────────────────────────────────────────────────────
def repartir(personas, mesas, fecha, semilla=7, intentos=300):
    """Asigna p['mesa'] a quien no la tiene, respetando fijos y reservas."""
    rnd = random.Random(semilla)
    por_n = {m["n"]: m for m in mesas}
    # Fijos (decisión humana) y reservas (grupo en su mesa unida).
    reserva_de = {m["reservada"]: m["n"] for m in mesas if m["reservada"]}
    for p in personas:
        if p["mesa"] in (None, "") and not p["fijo"] and clave_unidad(p) in reserva_de:
            p["mesa"] = reserva_de[clave_unidad(p)]
            p["fijo"] = True
    fijos, sueltos = {}, {}
    for p in personas:
        if p["mesa"] == "-":
            continue
        if p["fijo"] and p["mesa"] not in (None, ""):
            if p["mesa"] not in por_n:
                falla(f"{p['nombre']} está fijo en la mesa {p['mesa']}, que no existe en el Salón.")
            fijos.setdefault((p["mesa"], clave_unidad(p)), []).append(p)
        else:
            sueltos.setdefault(clave_unidad(p), []).append(p)
    unidades = [Unidad(k, ps, fecha) for k, ps in sueltos.items()]
    for u in unidades:
        previas = [p.get("previa") for p in u.personas if p.get("previa")]
        u.previa = max(set(previas), key=previas.count) if previas else None
    ancla = {n: [] for n in por_n}           # unidades fijas por mesa (entran al puntaje)
    for (n, k), ps in fijos.items():
        ancla[n].append(Unidad(k, ps, fecha))
    cap = {n: por_n[n]["puestos"] - sum(u.n for u in ancla[n]) for n in por_n}
    for n, c in cap.items():
        if c < 0:
            aviso(f"⚠ La mesa {n} tiene {-c} persona(s) más que sus puestos (fijadas a mano).")
    # Las mesas reservadas son del grupo: el cálculo no sienta a nadie más ahí.
    abiertas = [n for n in por_n if not por_n[n]["reservada"]]
    total = sum(u.n for u in unidades)
    if total > sum(max(cap[n], 0) for n in abiertas):
        falla(f"{total} viajeros por sentar y {sum(max(cap[n], 0) for n in abiertas)} puestos libres. "
              "Suma mesas (--sumar-mesas o una fila nueva en la hoja Salón).")
    if unidades and not abiertas:
        falla("Todas las mesas están reservadas y quedan viajeros sin mesa. Suma una mesa en el Salón.")
    grande = [u for u in unidades if u.n > max(cap[n] for n in abiertas)] if unidades else []
    if grande:
        falla(f"{grande[0].nombre} ({grande[0].n}) no cabe en ninguna mesa libre: resérvale mesas "
              f"unidas (--reservar {grande[0].clave}:2 o en la hoja Salón).")

    c_af, c_sec = {}, {}

    def memo(cache, f, a, b):
        k = (id(a), id(b))
        if k not in cache:
            cache[k] = cache[(id(b), id(a))] = f(a, b)
        return cache[k]

    def score_mesa(n, us):
        todos = ancla[n] + us
        if not todos:
            return 0.0
        # Cada unidad suma la RAÍZ de lo que comparte con su mesa. Rendimiento
        # decreciente a propósito: pasar de nada a un vínculo vale más que pasar
        # de tres a cuatro. Sumando parejo, el cálculo juntaba todos los vínculos
        # en pocas mesas y dejaba más mesas sin nada de qué hablar que la
        # distribución hecha a mano del 21-09. "Uniforme" es esto: que a cada
        # familia le toque alguien con quien compartió algo.
        s = sum(PESO_VINCULO * sum(memo(c_af, afinidad, u, v) for v in todos if v is not u) ** 0.5
                for u in todos)
        s += sum(memo(c_sec, secundario, a, b) for a, b in combinations(todos, 2))
        s += ESTABILIDAD * sum(1 for u in us if u.previa == n)
        vacios = por_n[n]["puestos"] - sum(u.n for u in todos)
        return s - CASTIGO_MESA - CASTIGO_VACIO * max(vacios, 0) ** 2

    def libre(n, conj):
        return cap[n] - sum(u.n for u in conj)

    mejor, mejor_s = None, float("-inf")
    for _ in range(intentos):
        conj = {n: [] for n in abiertas}
        orden = sorted(unidades, key=lambda u: (-u.n, rnd.random()))
        ok = True
        for u in orden:
            opciones = [n for n in abiertas if libre(n, conj[n]) >= u.n]
            if not opciones:
                ok = False
                break
            gan = {n: score_mesa(n, conj[n] + [u]) - score_mesa(n, conj[n]) for n in opciones}
            top = max(gan.values())
            n = rnd.choice([n for n in opciones if gan[n] >= top - 1e-9])
            conj[n].append(u)
        if not ok:
            continue
        # Búsqueda local: mover una unidad o intercambiar dos, mientras mejore.
        donde = {id(u): n for n, us in conj.items() for u in us}
        mejora = True
        while mejora:
            mejora = False
            us = unidades[:]
            rnd.shuffle(us)
            for u in us:
                a = donde[id(u)]
                base_a = score_mesa(a, conj[a])
                sin_u = [x for x in conj[a] if x is not u]
                for b in abiertas:
                    if b == a:
                        continue
                    base_b = score_mesa(b, conj[b])
                    if libre(b, conj[b]) >= u.n:
                        d = score_mesa(a, sin_u) + score_mesa(b, conj[b] + [u]) - base_a - base_b
                        if d > 1e-9:
                            conj[a], conj[b] = sin_u, conj[b] + [u]
                            donde[id(u)] = b
                            mejora = True
                            break
                    for v in conj[b]:
                        if libre(b, conj[b]) + v.n < u.n or libre(a, sin_u) < v.n:
                            continue
                        na = sin_u + [v]
                        nb = [x for x in conj[b] if x is not v] + [u]
                        d = score_mesa(a, na) + score_mesa(b, nb) - base_a - base_b
                        if d > 1e-9:
                            conj[a], conj[b] = na, nb
                            donde[id(u)], donde[id(v)] = b, a
                            mejora = True
                            break
                    if mejora:
                        break
                if mejora:
                    break
        s = sum(score_mesa(n, conj[n]) for n in abiertas)
        if s > mejor_s:
            mejor, mejor_s = {n: list(us) for n, us in conj.items()}, s
    if mejor is None:
        falla("No encontré una distribución que respete los puestos. Suma mesas o suelta algún fijo.")
    for n, us in mejor.items():
        for u in us:
            for p in u.personas:
                p["mesa"] = n
    return mejor_s


# ── Lo que se muestra de cada mesa ────────────────────────────────────────────
def unidades_en(personas, n):
    grupos = {}
    for p in personas:
        if p["mesa"] == n:
            grupos.setdefault(clave_unidad(p), []).append(p)
    return grupos


def titulo_mesa(m, personas):
    g = unidades_en(personas, m["n"])
    if m.get("reservada"):
        return m["reservada"].title()
    orden = sorted(g.items(), key=lambda kv: (-len(kv[1]), orden_hab(kv[1][0]["hab"])))
    return " · ".join(nombre_unidad(k, ps) for k, ps in orden)


def nota_mesa(m, personas, fecha):
    """Por qué están juntos, en una línea — el 'Los tres hicieron…' del 21-09."""
    g = unidades_en(personas, m["n"])
    total = sum(len(v) for v in g.values())
    if m.get("reservada"):
        return f"Grupo completo · {total} personas. Llegan juntos." if m["fisicas"] > 1 else \
               f"Mesa propia · {total} personas."
    us = [Unidad(k, ps, fecha) for k, ps in g.items()]
    comp = {}
    for a, b in combinations(us, 2):
        for k in a.exps:
            if k in b.exps:
                comp.setdefault(k, set()).update({a.nombre, b.nombre})
    if comp:
        def y(xs):
            return ", ".join(xs[:-1]) + " y " + xs[-1] if len(xs) > 1 else xs[0]
        # Lo que más gente comparte primero; a igualdad, lo más reciente. Las
        # exploraciones de las MISMAS personas van juntas en una sola frase.
        top = sorted(comp.items(), key=lambda kv: (-len(kv[1]), -int(kv[0][2].replace("-", ""))))[:3]
        frases = {}
        for (n, t, d), quienes in top:
            nombre = next((e[0] for u in us for p in u.personas for e in p["exps"]
                           if norm_key(e[0]) == n and e[1] == t and e[2] == d), n)
            frases.setdefault(tuple(sorted(quienes)), []).append(f"{nombre} ({t} {d[8:10]}-{d[5:7]})")
        partes = []
        for qs, exps in frases.items():
            quien = "Todos" if len(qs) == len(us) and len(us) > 2 else y(list(qs))
            partes.append(f"{quien}: {y(exps)}")
        return ". ".join(partes) + "."
    idiomas = sorted({IDIOMA_NOMBRE.get(u.idioma, u.idioma) for u in us})
    edades = [e for u in us for e in u.edades if e >= 18]
    rango = f" · {min(edades)} a {max(edades)} años" if edades else ""
    lenguas = ", ".join(idiomas[:-1]) + " y " + idiomas[-1] if len(idiomas) > 1 else "".join(idiomas)
    return f"Sin exploración compartida · {lenguas}{rango}." if us else ""


# ── Planilla ──────────────────────────────────────────────────────────────────
H_HAB = ["Hab", "Viajero", "Mesa", "Grupo", "Fijo", "Propuesta", "pid"]
H_SALON = ["Mesa", "Posición", "Puestos", "Mesas físicas", "Reservada para", "Ocupados", "Nota"]


def escribir_planilla(ruta, fecha, personas, mesas):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.properties import PageSetupProperties

    tinta, suave, ambar = "23201C", "7A7267", "B5832B"
    f_tit = Font(name="Calibri", size=14, bold=True, color=tinta)
    f_cab = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    f_sua = Font(name="Calibri", size=9, italic=True, color=suave)
    f_mesa = Font(name="Calibri", size=12, bold=True, color=ambar)
    fill_cab = PatternFill("solid", fgColor=tinta)
    fill_grupo = PatternFill("solid", fgColor="FBF3E2")
    linea = Border(bottom=Side(style="hair", color="D8D2C6"))
    por_n = {m["n"]: m for m in mesas}
    orden_fisico = [m["n"] for m in sorted(mesas, key=lambda m: m["pos"])]
    dia = fecha.strftime("%d-%m-%Y")

    wb = Workbook()
    # 1 · Por habitación: la que consulta el anfitrión y la que se edita.
    ws = wb.active
    ws.title = "Por habitación"
    ws["A1"] = f"Quincho · {dia} · {sum(1 for p in personas if p['mesa'] not in (None, '', '-'))} viajeros"
    ws["A1"].font = f_tit
    ws["A2"] = ("Salón de izquierda a derecha: " + " · ".join(str(n) for n in orden_fisico) +
                ".   Para mover a alguien, cambia su Mesa. «-» = no va. Fijo «x» = no lo muevas al recalcular.")
    ws["A2"].font = f_sua
    for i, h in enumerate(H_HAB, 1):
        c = ws.cell(row=4, column=i, value=h)
        c.font, c.fill = f_cab, fill_cab
    fila = 5
    for p in sorted(personas, key=lambda p: (orden_hab(p["hab"]), p["nombre"])):
        mesa = p["mesa"] if p["mesa"] not in (None, "") else ""
        vals = [p["hab"], p["nombre"], mesa, p["grupo"], "x" if p["fijo"] and not
                (por_n.get(mesa, {}).get("reservada") == clave_unidad(p)) else "",
                mesa, p["pid"]]
        for i, v in enumerate(vals, 1):
            c = ws.cell(row=fila, column=i, value=v)
            c.border = linea
        ws.cell(row=fila, column=3).font = f_mesa
        ws.cell(row=fila, column=3).alignment = Alignment(horizontal="center")
        ws.cell(row=fila, column=5).alignment = Alignment(horizontal="center")
        ws.cell(row=fila, column=6).font = Font(size=9, color="B8B0A3")
        ws.cell(row=fila, column=6).alignment = Alignment(horizontal="center")
        fila += 1
    for col, w in zip("ABCDEFG", (8, 38, 7, 16, 6, 10, 12)):
        ws.column_dimensions[col].width = w
    ws.column_dimensions["G"].hidden = True
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:F{fila - 1}"
    ws.print_title_rows = "4:4"
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0

    # 2 · Por mesa: el plano para imprimir. Se regenera; no se edita acá.
    wp = wb.create_sheet("Por mesa")
    wp["A1"] = f"Quincho · plano de mesas · {dia}"
    wp["A1"].font = f_tit
    wp["A2"] = "Se regenera con quincho.py a partir de «Por habitación»: los cambios se hacen allá."
    wp["A2"].font = f_sua
    fila = 4
    for n in orden_fisico:
        m = por_n[n]
        g = unidades_en(personas, n)
        total = sum(len(v) for v in g.values())
        c = wp.cell(row=fila, column=1, value=f"Mesa {n}")
        c.font = f_mesa
        wp.cell(row=fila, column=2, value=f"{titulo_mesa(m, personas)}  ·  {total}/{m['puestos']}").font = \
            Font(size=11, bold=True, color=tinta)
        if m["reservada"]:
            for col in (1, 2):
                wp.cell(row=fila, column=col).fill = fill_grupo
        fila += 1
        for k, ps in sorted(g.items(), key=lambda kv: orden_hab(kv[1][0]["hab"])):
            for hab in sorted({p["hab"] for p in ps}, key=orden_hab):
                wp.cell(row=fila, column=1, value=hab).font = Font(bold=True, color=ambar)
                wp.cell(row=fila, column=2, value=" · ".join(p["nombre"] for p in ps if p["hab"] == hab))
                fila += 1
        wp.cell(row=fila, column=2, value=m.get("nota") or "").font = f_sua
        fila += 2
    wp.column_dimensions["A"].width = 10
    wp.column_dimensions["B"].width = 90
    wp.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    wp.page_setup.fitToWidth, wp.page_setup.fitToHeight = 1, 0

    # 3 · Salón: las mesas que hay. Se editan acá y el cálculo las respeta.
    wsa = wb.create_sheet("Salón")
    wsa["A1"] = "Salón · las mesas disponibles"
    wsa["A1"].font = f_tit
    wsa["A2"] = ("Suma una fila para una mesa nueva. «Reservada para» = el grupo va solo en esa mesa; "
                 "«Mesas físicas» 2 o más = mesas unidas.")
    wsa["A2"].font = f_sua
    for i, h in enumerate(H_SALON, 1):
        c = wsa.cell(row=4, column=i, value=h)
        c.font, c.fill = f_cab, fill_cab
    for r, m in enumerate(sorted(mesas, key=lambda m: m["pos"]), 5):
        vals = [m["n"], m["pos"], m["puestos"], m["fisicas"], m["reservada"].title() if m["reservada"] else "",
                f"=COUNTIF('Por habitación'!C:C,A{r})", m.get("nota") or ""]
        for i, v in enumerate(vals, 1):
            wsa.cell(row=r, column=i, value=v).border = linea
    for col, w in zip("ABCDEFG", (7, 9, 8, 13, 18, 10, 70)):
        wsa.column_dimensions[col].width = w
    wsa.freeze_panes = "A5"

    # 4 · _datos (oculta): lo que el cálculo necesita para volver a correr.
    wd = wb.create_sheet("_datos")
    wd.sheet_state = "hidden"
    wd["A1"], wd["B1"] = "meta", json.dumps({"fecha": fecha.isoformat(), "version": 1})
    for r, p in enumerate(personas, 2):
        wd.cell(row=r, column=1, value=p["pid"])
        wd.cell(row=r, column=2, value=json.dumps(
            {k: p[k] for k in ("edad", "nac", "grupo", "exps")}, ensure_ascii=False))
    try:
        wb.save(ruta)
    except PermissionError:
        falla(f"No pude escribir {ruta}: ¿está abierto en Excel? Ciérralo y vuelve a correr.")


def leer_planilla(ruta):
    from openpyxl import load_workbook
    if not os.path.exists(ruta):
        falla(f"No existe {ruta}.")
    wb = load_workbook(ruta, data_only=False)
    datos, fecha = {}, None
    if "_datos" in wb.sheetnames:
        for pid, js in wb["_datos"].iter_rows(min_row=1, max_col=2, values_only=True):
            if pid == "meta":
                fecha = datetime.date.fromisoformat(json.loads(js)["fecha"])
            elif pid:
                datos[str(pid)] = json.loads(js)
    if not fecha:
        m = re.search(r"(\d{4}-\d{2}-\d{2})", os.path.basename(ruta))
        fecha = datetime.date.fromisoformat(m.group(1)) if m else datetime.date.today()

    def entero(v):
        try:
            return int(float(str(v).strip()))
        except (TypeError, ValueError):
            return None

    mesas = []
    for fila in wb["Salón"].iter_rows(min_row=5, max_col=7, values_only=True):
        n = entero(fila[0])
        if n is None:
            continue
        mesas.append({"n": n, "pos": entero(fila[1]) or n, "puestos": entero(fila[2]) or PUESTOS_BASE,
                      "fisicas": entero(fila[3]) or 1,
                      "reservada": str(fila[4] or "").strip().upper(), "nota": str(fila[6] or "")})
    if len({m["n"] for m in mesas}) != len(mesas):
        falla("Hay números de mesa repetidos en la hoja Salón.")
    personas = []
    for fila in wb["Por habitación"].iter_rows(min_row=5, max_col=7, values_only=True):
        hab, nombre, mesa, grupo, fijo, propuesta, pid = (list(fila) + [None] * 7)[:7]
        if not nombre:
            continue
        nombre = str(nombre).strip()
        pid = str(pid).strip() if pid else pid_de(nombre)
        d = datos.get(pid, {})
        txt_mesa = str(mesa).strip().lower() if mesa is not None else ""
        if txt_mesa in NO_VA:
            m = "-"
        else:
            m = entero(mesa)
        prop = entero(propuesta)
        # Fijo si lo marcaron, si la mesa difiere de la propuesta (alguien la
        # cambió), o si es una fila nueva que ya trae mesa.
        es_fijo = (str(fijo or "").strip().lower() in SI) or \
                  (m not in (None, "-") and (prop is None or m != prop))
        personas.append({"pid": pid, "nombre": nombre, "hab": str(hab or "").strip() or "SIN HAB",
                         "edad": d.get("edad"), "nac": d.get("nac", ""),
                         "grupo": str(grupo or d.get("grupo") or "").strip(),
                         "exps": [tuple(e) for e in d.get("exps", [])],
                         "mesa": m, "fijo": es_fijo, "previa": prop})
    return fecha, personas, mesas


# ── Carteles ──────────────────────────────────────────────────────────────────
def logo_data_uri():
    ruta = os.path.join(REPO, "assets", "logo-explora-mascara.png")   # versionado; el de la raíz lo ignora *.png
    with open(ruta, "rb") as f:
        return "data:image/png;base64," + base64.b64encode(f.read()).decode()


def escribir_carteles(ruta, fecha, personas, mesas):
    """A4 vertical, se dobla al medio; la mitad de arriba va invertida.

    El logo es una MÁSCARA: la forma sale del PNG y el color de una variable CSS,
    así se imprime del color que se elija sobre papel blanco. Todo el texto es
    editable con un clic, y lo que se corrige en una mitad se copia a la otra.
    """
    e = html.escape

    def cara(m):
        g = unidades_en(personas, m["n"])
        total = sum(len(v) for v in g.values())
        grande = m["fisicas"] > 1
        if grande:
            habs = " · ".join(sorted({p["hab"] for ps in g.values() for p in ps}, key=orden_hab))
            cuerpo = (f'<div class="grupo">Grupo completo · {total} personas</div>'
                      f'<div class="habs">habitaciones {e(habs)}</div>')
        else:
            filas = []
            for k, ps in sorted(g.items(), key=lambda kv: orden_hab(kv[1][0]["hab"])):
                for hab in sorted({p["hab"] for p in ps}, key=orden_hab):
                    noms = " · ".join(e(p["nombre"]) for p in ps if p["hab"] == hab)
                    filas.append(f'<div class="u"><span class="hab">{e(hab)}</span>'
                                 f'<span class="nom">{noms}</span></div>')
            cuerpo = f'<div class="lista">{"".join(filas)}</div>'
        cls = "cara grande" if grande else "cara"
        return (f'<div class="{cls}" data-mesa="{m["n"]}">'
                f'<div class="cifra">{m["n"]}</div><div class="rotulo">MESA</div>'
                f'<h1>{e(titulo_mesa(m, personas))}</h1>{cuerpo}'
                f'<div class="logo" role="img" aria-label="explora"></div></div>')

    usadas = [m for m in sorted(mesas, key=lambda m: m["n"]) if unidades_en(personas, m["n"])]
    paginas = "".join(
        f'<section class="hoja"><div class="mitad invertida">{cara(m)}</div>'
        f'<div class="pliegue"></div><div class="mitad">{cara(m)}</div></section>'
        for m in usadas)
    doc = PLANTILLA_CARTELES.replace("__PAGINAS__", paginas).replace("__LOGO__", logo_data_uri()) \
        .replace("__FECHA__", fecha.strftime("%d-%m-%Y")).replace("__N__", str(len(usadas)))
    with open(ruta, "w", encoding="utf-8") as f:
        f.write(doc)


PLANTILLA_CARTELES = """<!DOCTYPE html>
<html lang="es"><head><meta charset="utf-8"><title>Carteles de mesa · Quincho __FECHA__</title>
<link rel="preconnect" href="https://fonts.googleapis.com">
<link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
<link href="https://fonts.googleapis.com/css2?family=Source+Sans+3:wght@400;600;700&family=Cormorant+Garamond:wght@600&display=swap" rel="stylesheet">
<style>
@page { size: A4; margin: 0; }
:root{--tinta:#23201C;--suave:#7A7267;--linea:#D8D2C6;--ambar:#B5832B;--papel:#FFFFFF;--logo:#23201C;}
*{box-sizing:border-box}
body{margin:0;background:#E9E6E0;color:var(--tinta);
 font-family:"Source Sans 3",system-ui,sans-serif;-webkit-print-color-adjust:exact;print-color-adjust:exact}

.barra{position:sticky;top:0;z-index:5;display:flex;flex-wrap:wrap;gap:14px;align-items:center;
 padding:10px 16px;background:var(--tinta);color:#E8E3D9;font-size:14px}
.barra b{font-weight:600}
.barra label{display:flex;align-items:center;gap:6px}
.barra input[type=color]{width:34px;height:24px;border:0;padding:0;background:none;cursor:pointer}
.barra .muestra{width:20px;height:20px;border-radius:50%;border:1px solid #fff5;cursor:pointer}
.barra button{margin-left:auto;background:var(--ambar);color:#fff;border:0;border-radius:4px;
 padding:7px 16px;font:inherit;font-weight:600;cursor:pointer}
.barra .pista{opacity:.7;font-size:12.5px}

.hoja{width:210mm;height:297mm;display:flex;flex-direction:column;margin:12px auto;
 break-after:page;background:var(--papel);overflow:hidden;box-shadow:0 1px 6px #0002}
.mitad{flex:1 1 50%;display:flex;align-items:center;justify-content:center;padding:10mm 14mm;min-height:0}
.invertida{transform:rotate(180deg)}
.pliegue{height:0;border-top:0.4pt dashed var(--linea)}
.cara{width:100%;text-align:center}
.cara [contenteditable]:hover{outline:1px dashed #B5832B66;outline-offset:2px}
.cara [contenteditable]:focus{outline:1px solid var(--ambar);outline-offset:2px}

/* Cifras de caja alta: las de estilo antiguo de Cormorant (3, 4, 5, 7, 9) bajan
   de la línea y se montaban sobre el rótulo MESA. */
.cifra{font-family:"Cormorant Garamond",Georgia,serif;font-size:120pt;line-height:.82;color:var(--ambar);margin:0;
 font-variant-numeric:lining-nums;font-feature-settings:"lnum" 1}
.rotulo{font-size:9pt;letter-spacing:.34em;color:var(--suave);margin:2mm 0 1mm;padding-left:.34em}
h1{font-family:"Cormorant Garamond",Georgia,serif;font-weight:600;font-size:30pt;line-height:1.08;
 margin:0 0 4mm;padding-bottom:3mm;border-bottom:1pt solid var(--ambar)}
.cara.grande h1{font-size:52pt;letter-spacing:.01em}
.grupo{font-size:14pt;color:var(--suave)}
.habs{font-size:10.5pt;color:var(--ambar);font-weight:600;margin-top:2.5mm;letter-spacing:.02em}

/* El logo es la forma del PNG rellena con --logo: cualquier color, nítido. */
.logo{width:15mm;height:15mm;margin:11mm auto 0;background:var(--logo);
 -webkit-mask:url(__LOGO__) center/contain no-repeat;mask:url(__LOGO__) center/contain no-repeat}
.cara.grande .logo{width:17mm;height:17mm;margin-top:14mm}

.lista{display:inline-block;text-align:left}
.u{display:flex;gap:4mm;padding:1.1mm 0;font-size:11.5pt;line-height:1.25}
.hab{flex:0 0 17mm;color:var(--ambar);font-weight:700;text-align:right}
.nom{flex:1}

@media print{
  body{background:none}
  .barra{display:none}
  .hoja{margin:0;box-shadow:none}
  .cara [contenteditable]{outline:none!important}
}
</style></head><body>
<div class="barra">
  <b>Carteles · Quincho __FECHA__ · __N__ mesas</b>
  <label>Logo <input type="color" id="c-logo" value="#23201C"></label>
  <span class="muestra" style="background:#23201C" data-c="#23201C" title="Tinta"></span>
  <span class="muestra" style="background:#B5832B" data-c="#B5832B" title="Ámbar"></span>
  <span class="muestra" style="background:#7A7267" data-c="#7A7267" title="Gris cálido"></span>
  <span class="pista">Un clic sobre cualquier texto lo corrige. Imprimir: A4, márgenes «Ninguno», con gráficos de fondo.</span>
  <button type="button" onclick="window.print()">Imprimir</button>
</div>
__PAGINAS__
<script>
// Color del logo: una variable CSS para todos los carteles.
const input = document.getElementById('c-logo');
const pintar = c => { document.documentElement.style.setProperty('--logo', c); input.value = c; };
input.addEventListener('input', e => pintar(e.target.value));
document.querySelectorAll('.muestra').forEach(m => m.onclick = () => pintar(m.dataset.c));

// Todo el texto se corrige con un clic, y la otra mitad del mismo cartel se
// copia sola: al doblar, las dos caras tienen que decir lo mismo.
document.querySelectorAll('.cara .cifra, .cara h1, .cara .grupo, .cara .habs, .cara .hab, .cara .nom')
  .forEach(el => { el.contentEditable = 'true'; el.spellcheck = false; });
document.addEventListener('input', e => {
  const cara = e.target.closest('.cara');
  if (!cara) return;
  const hoja = cara.closest('.hoja');
  const gemela = [...hoja.querySelectorAll('.cara')].find(c => c !== cara);
  if (gemela) gemela.innerHTML = cara.innerHTML;
});
</script>
</body></html>"""


# ── Consola ───────────────────────────────────────────────────────────────────
def resumen(fecha, personas, mesas):
    por_n = {m["n"]: m for m in mesas}
    sentados = [p for p in personas if p["mesa"] not in (None, "", "-")]
    total_af = 0.0
    print(f"\n=== QUINCHO {fecha.strftime('%d-%m-%Y')} · {len(sentados)} viajeros · "
          f"{sum(1 for m in mesas if unidades_en(personas, m['n']))} mesas en uso ===")
    print("Salón, de izquierda a derecha: " +
          " · ".join(str(m["n"]) for m in sorted(mesas, key=lambda m: m["pos"])))
    for m in sorted(mesas, key=lambda m: m["pos"]):
        g = unidades_en(personas, m["n"])
        total = sum(len(v) for v in g.values())
        us = [Unidad(k, ps, fecha) for k, ps in g.items()]
        af = sum(afinidad(a, b) for a, b in combinations(us, 2))
        total_af += af
        marca = " [unida]" if m["fisicas"] > 1 else (" [reservada]" if m["reservada"] else "")
        print(f"\nMesa {m['n']:>2}  {total}/{m['puestos']}{marca}  {titulo_mesa(m, personas) or '(libre)'}")
        for k, ps in sorted(g.items(), key=lambda kv: orden_hab(kv[1][0]["hab"])):
            fij = " (fijo)" if any(p["fijo"] for p in ps) and not m["reservada"] else ""
            print(f"    {', '.join(sorted({p['hab'] for p in ps}, key=orden_hab)):8s} "
                  f"{' · '.join(p['nombre'] for p in ps)}{fij}")
        if m.get("nota"):
            print(f"    → {m['nota']}")
    fuera = [p for p in personas if p["mesa"] == "-"]
    if fuera:
        print(f"\nNo van: {', '.join(p['nombre'] for p in fuera)}")
    print(f"\nAfinidad por exploraciones compartidas (ponderada por recencia): {total_af:.1f}")
    llenas = [m for m in mesas if sum(len(v) for v in unidades_en(personas, m['n']).values()) > m["puestos"]]
    for m in llenas:
        print(f"⚠ Mesa {m['n']} excede sus puestos.")


def rutas(base_dir, fecha):
    stem = os.path.join(base_dir, f"quincho-{fecha.isoformat()}")
    return stem + ".xlsx", stem + "-carteles.html"


def completar_notas(fecha, personas, mesas):
    for m in mesas:
        m["nota"] = nota_mesa(m, personas, fecha) if unidades_en(personas, m["n"]) else ""


# ── Comandos ──────────────────────────────────────────────────────────────────
def cmd_proponer(a):
    fecha = datetime.date.fromisoformat(a.fecha) if a.fecha else datetime.date.today()
    if a.roster:
        with open(a.roster, encoding="utf-8") as f:
            doc = json.load(f)
    else:
        aviso("Leyendo /viajeros/current de Firebase…")
        doc = fb("GET", "viajeros/current")
    if not doc or not doc.get("habs"):
        falla("El roster vino vacío.")
    personas = personas_del_roster(doc, fecha, a.excluir or [])
    if not personas:
        falla("No hay viajeros en el hotel esa noche.")
    # Reservas: las pedidas, las de mesa propia y las de los grupos que no
    # caben en una mesa. Lo explícito gana a lo automático.
    tam = {}
    for p in personas:
        if p["mesa"] != "-":
            tam[clave_unidad(p)] = tam.get(clave_unidad(p), 0) + 1
    reservas = {}
    for k, n in tam.items():
        if n > a.puestos:
            reservas[k] = mesas_para(n, a.puestos, a.puntas)
    for g in a.mesa_propia or []:
        k = g.strip().upper()
        if k not in tam:
            falla(f"--mesa-propia {g}: no hay ningún grupo con ese nombre esa noche.")
        reservas[k] = mesas_para(tam[k], a.puestos, a.puntas)
    for r in a.reservar or []:
        g, _, k = r.rpartition(":")
        if not g or not k.isdigit():
            falla(f"--reservar {r}: se escribe GRUPO:MESAS, por ejemplo PATIL:2.")
        if g.strip().upper() not in tam:
            falla(f"--reservar {r}: no hay ningún grupo {g} esa noche.")
        reservas[g.strip().upper()] = int(k)
    for k, n in reservas.items():
        cap = n * a.puestos + (a.puntas if n > 1 else 0)
        if tam[k] > cap:
            falla(f"{k} son {tam[k]} y {n} mesa(s) dan {cap} puestos.")
        aviso(f"{k.title()}: {tam[k]} personas → {n} mesa(s){' unidas' if n > 1 else ''}.")
    mesas = armar_salon(a.mesas + a.sumar_mesas, a.puestos, a.puntas, reservas)
    repartir(personas, mesas, fecha, a.semilla, a.intentos)
    completar_notas(fecha, personas, mesas)
    os.makedirs(a.salida, exist_ok=True)
    xlsx, carteles = rutas(a.salida, fecha)
    escribir_planilla(xlsx, fecha, personas, mesas)
    escribir_carteles(carteles, fecha, personas, mesas)
    resumen(fecha, personas, mesas)
    print(f"\n→ {xlsx}\n→ {carteles}")


def cmd_rehacer(a):
    fecha, personas, mesas = leer_planilla(a.archivo)
    fijos = sum(1 for p in personas if p["fijo"])
    antes = {p["pid"]: p["previa"] for p in personas}
    for p in personas:
        if not p["fijo"] and p["mesa"] != "-":
            p["mesa"] = None
        if a.desde_cero:
            p["previa"] = None
    aviso(f"{fijos} viajero(s) fijos se respetan; el resto se vuelve a repartir"
          + (" desde cero." if a.desde_cero else ", moviendo lo mínimo."))
    repartir(personas, mesas, fecha, a.semilla, a.intentos)
    completar_notas(fecha, personas, mesas)
    # Qué carteles cambiaron: la mesa que dejó y la que ganó cada movido.
    tocadas = set()
    for p in personas:
        previa = antes.get(p["pid"])
        if p["mesa"] != previa:
            tocadas.update(x for x in (previa, p["mesa"]) if isinstance(x, int))
    shutil.copyfile(a.archivo, a.archivo.replace(".xlsx", ".anterior.xlsx"))
    escribir_planilla(a.archivo, fecha, personas, mesas)
    carteles = a.archivo.replace(".xlsx", "-carteles.html")
    escribir_carteles(carteles, fecha, personas, mesas)
    resumen(fecha, personas, mesas)
    if tocadas:
        print(f"\nCarteles a reimprimir: mesa {', '.join(str(n) for n in sorted(tocadas))}.")
    else:
        print("\nNingún cartel cambió.")
    print(f"\n→ {a.archivo}  (la versión previa quedó como .anterior.xlsx)\n→ {carteles}")


def cmd_carteles(a):
    fecha, personas, mesas = leer_planilla(a.archivo)
    carteles = a.archivo.replace(".xlsx", "-carteles.html")
    escribir_carteles(carteles, fecha, personas, mesas)
    print(f"→ {carteles}")


def doc_handbook(fecha, personas, mesas):
    """El doc de /quincho/current. `asig` va por pid: el mismo ancla que las notas."""
    completar_notas(fecha, personas, mesas)
    return {
        "fecha": fecha.isoformat(),
        "updatedAt": int(datetime.datetime.now(datetime.timezone.utc).timestamp() * 1000),
        "fuente": "quincho.py",
        "mesas": {str(m["n"]): {"n": m["n"], "pos": m["pos"], "puestos": m["puestos"],
                                "fisicas": m["fisicas"], "grupo": bool(m["reservada"]),
                                "titulo": titulo_mesa(m, personas), "nota": m.get("nota") or ""}
                  for m in mesas},
        "asig": {p["pid"]: {"m": p["mesa"], "hab": p["hab"], "nombre": p["nombre"],
                            "grupo": p["grupo"]}
                 for p in personas if p["mesa"] not in (None, "", "-")},
    }


def cmd_publicar(a):
    fecha, personas, mesas = leer_planilla(a.archivo)
    sin_mesa = [p["nombre"] for p in personas if p["mesa"] in (None, "")]
    if sin_mesa:
        falla(f"Hay viajeros sin mesa ({', '.join(sin_mesa[:3])}…): corre `rehacer` antes de publicar.")
    doc = doc_handbook(fecha, personas, mesas)
    print(f"[quincho] /{QUINCHO_PATH}: {len(doc['asig'])} viajeros en "
          f"{sum(1 for m in mesas if unidades_en(personas, m['n']))} mesas · fecha {doc['fecha']}")
    if a.debug:
        aviso("--debug: no se escribió nada.")
        return
    fb("PUT", QUINCHO_PATH, doc)
    aviso("Publicado. En el Handbook: PGO → Quincho.")


def cmd_bajar(a):
    from openpyxl import load_workbook
    vivo = fb("GET", QUINCHO_PATH)
    if not vivo or not vivo.get("asig"):
        falla("No hay un quincho publicado en el Handbook.")
    wb = load_workbook(a.archivo)
    ws = wb["Por habitación"]
    cambios = 0
    for fila in ws.iter_rows(min_row=5):
        nombre, pid = fila[1].value, fila[6].value
        if not nombre:
            continue
        pid = str(pid) if pid else pid_de(str(nombre))
        v = (vivo["asig"] or {}).get(pid)
        if v and str(v.get("m")) != str(fila[2].value):
            aviso(f"{nombre}: mesa {fila[2].value} → {v.get('m')}")
            fila[2].value = v.get("m")
            fila[4].value = "x"
            cambios += 1
    if not cambios:
        aviso("La planilla ya está igual que el Handbook.")
        return
    try:
        wb.save(a.archivo)
    except PermissionError:
        falla(f"No pude escribir {a.archivo}: ¿está abierto en Excel?")
    aviso(f"{cambios} cambio(s) traídos y marcados como Fijo. `carteles` los lleva a los carteles.")


def main():
    ap = argparse.ArgumentParser(description="Mesas del quincho: propone, respeta lo ajustado a mano, imprime y publica.")
    sub = ap.add_subparsers(dest="cmd", required=True)
    p = sub.add_parser("proponer", help="calcula desde el roster y escribe planilla + carteles")
    p.add_argument("--fecha")
    p.add_argument("--roster")
    p.add_argument("--mesas", type=int, default=MESAS_BASE)
    p.add_argument("--puestos", type=int, default=PUESTOS_BASE)
    p.add_argument("--puntas", type=int, default=PUNTAS)
    p.add_argument("--sumar-mesas", type=int, default=0)
    p.add_argument("--reservar", action="append")
    p.add_argument("--mesa-propia", action="append")
    p.add_argument("--excluir", action="append")
    p.add_argument("--salida", default=".")
    for s in (p, sub.add_parser("rehacer", help="recalcula respetando lo fijado en la planilla")):
        if s is not p:
            s.add_argument("archivo")
            s.add_argument("--desde-cero", action="store_true",
                           help="ignora las mesas actuales de quien no está fijo")
        s.add_argument("--semilla", type=int, default=7)
        s.add_argument("--intentos", type=int, default=300)
    sub.add_parser("carteles", help="regenera sólo los carteles").add_argument("archivo")
    q = sub.add_parser("publicar", help="sube la planilla al Handbook")
    q.add_argument("archivo")
    q.add_argument("--debug", action="store_true")
    sub.add_parser("bajar", help="trae a la planilla lo movido en el Handbook").add_argument("archivo")
    a = ap.parse_args()
    {"proponer": cmd_proponer, "rehacer": cmd_rehacer, "carteles": cmd_carteles,
     "publicar": cmd_publicar, "bajar": cmd_bajar}[a.cmd](a)


if __name__ == "__main__":
    main()
